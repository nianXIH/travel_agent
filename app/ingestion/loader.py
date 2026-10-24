from pathlib import Path

ENCODINGS = ("utf-8-sig", "utf-8", "gb18030")
TEXT_SUFFIXES = {".md", ".markdown", ".txt"}


class DocumentLoadError(RuntimeError):
    pass


def load_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return _with_title(path, _load_plain_text(path))
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".docx":
        return _load_docx(path)
    raise DocumentLoadError(f"unsupported file type: {path.suffix}")


def _load_plain_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ENCODINGS:
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise DocumentLoadError(f"cannot decode {path} with {ENCODINGS}")

    return text.replace("\x00", "").strip()


def _load_pdf(path: Path) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise DocumentLoadError(f"cannot read pdf {path}: {exc}") from exc

    text = text.replace("\x00", "").strip()
    if not text:
        raise DocumentLoadError(f"no extractable text in {path}")
    return _with_title(path, text)


def _load_docx(path: Path) -> str:
    from docx import Document

    try:
        document = Document(str(path))
    except Exception as exc:
        raise DocumentLoadError(f"cannot read docx {path}: {exc}") from exc

    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").strip()
        if style.startswith("Heading "):
            level = style.removeprefix("Heading ").strip()
            lines.append(f"{'#' * int(level)} {text}" if level.isdigit() else f"# {text}")
        else:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    text = "\n".join(lines).strip()
    if not text:
        raise DocumentLoadError(f"no extractable text in {path}")
    return _with_title(path, text)


def _with_title(path: Path, text: str) -> str:
    if any(line.lstrip().startswith("#") for line in text.splitlines()):
        return text
    return f"# {path.stem}\n\n{text}"
